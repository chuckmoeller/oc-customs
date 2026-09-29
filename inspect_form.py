import openpyxl

wb = openpyxl.load_workbook(/app/app/templates/Submission_Form_8.xlsx, data_only=False)
print(Sheet names:, wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f\n--- Sheet: {sname} ({ws.max_row} rows, {ws.max_column} cols) ---)
    for r in range(1, min(ws.max_row + 1, 35)):
        row_vals = [f{openpyxl.utils.get_column_letter(c)}{r}={repr(ws.cell(r, c).value)} for c in range(1, min(ws.max_column + 1, 15)) if ws.cell(r, c).value is not None]
        if row_vals:
            print(fRow {r}:  + , .join(row_vals[:8]))


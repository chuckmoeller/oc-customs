from datetime import datetime
import io
import logging
import os
import re
from typing import Any, Dict, List, Optional
import uuid
from fastapi import APIRouter, Depends, HTTPException, Request, Response, status
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth import assert_job_in_org, get_auth_context
from app.db import get_session
from app.schemas import AuthContext
from app.services.classify import classify_system
from app.services.decoders import decode_serial_date
from app.services.devices_service import list_devices_for_job
from app.services.jobs_service import get_job_by_id

logger = logging.getLogger("export")
router = APIRouter(prefix="/jobs", tags=["Export"])

GROUP_LABEL = {"hvac": "HVAC", "cold": "Cold Storage", "other": "Other"}
GROUP_RANK = {"hvac": 0, "cold": 1, "other": 2}

TEMPLATE_CANDIDATES = [
    "/app/app/templates/Submission_Form_8.xlsx",
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates", "Submission_Form_8.xlsx"),
    "/home/opc/stack/app/templates/Submission_Form_8.xlsx",
    "/home/opc/stack/Submission_Form_8.xlsx",
]

PANEL_SPECS = [
    {"num": 1,  "name_cell": "C9",  "left_col": 2,  "right_col": 6,  "start_row": 10, "end_row": 31},
    {"num": 2,  "name_cell": "I9",  "left_col": 8,  "right_col": 12, "start_row": 10, "end_row": 31},
    {"num": 3,  "name_cell": "O9",  "left_col": 14, "right_col": 18, "start_row": 10, "end_row": 31},
    {"num": 4,  "name_cell": "U9",  "left_col": 20, "right_col": 24, "start_row": 10, "end_row": 31},
    {"num": 5,  "name_cell": "C36", "left_col": 2,  "right_col": 6,  "start_row": 37, "end_row": 58},
    {"num": 6,  "name_cell": "I36", "left_col": 8,  "right_col": 12, "start_row": 37, "end_row": 58},
    {"num": 7,  "name_cell": "O36", "left_col": 14, "right_col": 18, "start_row": 37, "end_row": 58},
    {"num": 8,  "name_cell": "U36", "left_col": 20, "right_col": 24, "start_row": 37, "end_row": 58},
    {"num": 9,  "name_cell": "C62", "left_col": 2,  "right_col": 6,  "start_row": 63, "end_row": 84},
    {"num": 10, "name_cell": "I62", "left_col": 8,  "right_col": 12, "start_row": 63, "end_row": 84},
    {"num": 11, "name_cell": "O62", "left_col": 14, "right_col": 18, "start_row": 63, "end_row": 84},
    {"num": 12, "name_cell": "U62", "left_col": 20, "right_col": 24, "start_row": 63, "end_row": 84},
    {"num": 13, "name_cell": "C88", "left_col": 2,  "right_col": 6,  "start_row": 89, "end_row": 110},
    {"num": 14, "name_cell": "I88", "left_col": 8,  "right_col": 12, "start_row": 89, "end_row": 110},
    {"num": 15, "name_cell": "O88", "left_col": 14, "right_col": 18, "start_row": 89, "end_row": 110},
    {"num": 16, "name_cell": "U88", "left_col": 20, "right_col": 24, "start_row": 89, "end_row": 110},
]

MONTH_ROW_MAP = {
    "january": 33, "jan": 33, "1": 33, "01": 33,
    "february": 34, "feb": 34, "2": 34, "02": 34,
    "march": 35, "mar": 35, "3": 35, "03": 35,
    "april": 36, "apr": 36, "4": 36, "04": 36,
    "may": 37, "5": 37, "05": 37,
    "june": 38, "jun": 38, "6": 38, "06": 38,
    "july": 39, "jul": 39, "7": 39, "07": 39,
    "august": 40, "aug": 40, "8": 40, "08": 40,
    "september": 41, "sep": 41, "sept": 41, "9": 41, "09": 41,
    "october": 42, "oct": 42, "10": 42,
    "november": 43, "nov": 43, "11": 43,
    "december": 44, "dec": 44, "12": 44,
}


def clean_float(val: Any) -> Optional[float]:
    """Safely extracts a floating-point number from string or numeric input."""
    if val is None:
        return None
    try:
        s = re.sub(r"[^\d.]", "", str(val))
        return float(s) if s else None
    except (ValueError, TypeError):
        return None


def clean_int(val: Any) -> Optional[int]:
    """Safely extracts an integer from string or numeric input."""
    if val is None:
        return None
    try:
        s = re.sub(r"[^\d]", "", str(val))
        return int(s) if s else None
    except (ValueError, TypeError):
        return None


def calc_age(mfg_year: Optional[int], mfg_date: Optional[str], serial: Optional[str], mfg: Optional[str]) -> str:
    """Calculates equipment age from manufacture year or serial number."""
    current_year = datetime.now().year
    if mfg_year and 1950 <= mfg_year <= current_year:
        return str(current_year - mfg_year)

    decoded = decode_serial_date(serial or "", mfg)
    if decoded and 1950 <= decoded <= current_year:
        return str(current_year - decoded)

    if mfg_date:
        m = re.search(r"\b(19[8-9]\d|20[0-3]\d)\b", mfg_date)
        if m:
            yr = int(m.group(1))
            if 1950 <= yr <= current_year:
                return str(current_year - yr)

    return ""


def csv_cell(val: Any) -> str:
    """Formats cell conforming to RFC-4180 with flattened newlines."""
    if val is None:
        return ""
    s = str(val)
    s = re.sub(r"\s*[\r\n]+\s*", "; ", s)
    s = re.sub(r"[ \t]{2,}", " ", s).strip()
    if '"' in s or "," in s:
        escaped = s.replace('"', '""')
        return f'"{escaped}"'
    return s


def format_equipment_item(item: Any, idx: int) -> str:
    """Extracts VLM nameplate data (Manufacturer, Model, Serial, Equipment Type) into an equipment label string."""
    analysis = getattr(item, "analysis_result", None) or {}
    mfg = (getattr(item, "manufacturer", None) or analysis.get("manufacturer") or "").strip()
    model = (getattr(item, "model_number", None) or analysis.get("model_number") or "").strip()
    serial = (getattr(item, "serial_number", None) or analysis.get("serial_number") or "").strip()
    eq_type = (getattr(item, "equipment_type", None) or analysis.get("equipment_type") or "").strip()
    dev_name = (getattr(item, "device_name", None) or getattr(item, "nameplate_label", None) or getattr(item, "name", None) or "").strip()

    parts = [p for p in [eq_type, mfg, model] if p]
    if parts:
        label = " - ".join(parts)
        if serial:
            label += f" (SN: {serial})"
    elif dev_name and dev_name.lower() not in ["scanned device", "none"]:
        label = dev_name
        if serial:
            label += f" (SN: {serial})"
    else:
        label = f"Equipment {idx + 1}"
    return label.strip()


def build_equipment_csv(devices: List[Any]) -> str:
    """Generates the standard 27-column Site Hunter equipment inventory CSV."""
    equipment_columns = [
        ("Name", lambda d, c: getattr(d, "device_name", None) or getattr(d, "name", None) or (d.get("name") if isinstance(d, dict) else "")),
        ("Type", lambda d, c: c.get("system")),
        ("Category", lambda d, c: GROUP_LABEL.get(c.get("group"), c.get("group"))),
        ("Qty", lambda d, c: getattr(d, "quantity", 1) or 1),
        ("Manufacturer", lambda d, c: getattr(d, "manufacturer", None) or (d.get("manufacturer") if isinstance(d, dict) else "")),
        ("Model", lambda d, c: getattr(d, "model_number", None) or (d.get("model_number") if isinstance(d, dict) else "")),
        ("Serial", lambda d, c: getattr(d, "serial_number", None) or (d.get("serial_number") if isinstance(d, dict) else "")),
        ("Age (yrs)", lambda d, c: calc_age(getattr(d, "mfg_year", None), getattr(d, "mfg_date", None), getattr(d, "serial_number", None), getattr(d, "manufacturer", None))),
        ("Mfg Date", lambda d, c: getattr(d, "mfg_date", None) or (str(getattr(d, "mfg_year", "")) if getattr(d, "mfg_year", None) else "")),
        ("Tonnage", lambda d, c: getattr(d, "tonnage", None) or (d.get("tonnage") if isinstance(d, dict) else "")),
        ("BTU", lambda d, c: getattr(d, "btu", None) or ""),
        ("Voltage", lambda d, c: getattr(d, "voltage", None) or (d.get("voltage") if isinstance(d, dict) else "")),
        ("Compressor PH", lambda d, c: getattr(d, "compressor_ph", None) or ""),
        ("Fan PH", lambda d, c: getattr(d, "fan_phases", None) or getattr(d, "fan_ph", None) or getattr(d, "compressor_ph", None) or (d.get("fan_phases") if isinstance(d, dict) else (d.get("fan_ph") if isinstance(d, dict) else (d.get("compressor_ph") if isinstance(d, dict) else (d.get("phase") if isinstance(d, dict) else "")))) or ""),
        ("Refrigerant", lambda d, c: getattr(d, "refrigerant_type", None) or (d.get("refrigerant") if isinstance(d, dict) else "")),
        ("MCA", lambda d, c: getattr(d, "mca", None) or ""),
        ("MOCP", lambda d, c: getattr(d, "mocp", None) or ""),
        ("Compressor RLA", lambda d, c: getattr(d, "compressor_rla", None) or ""),
        ("Compressor LRA", lambda d, c: getattr(d, "compressor_lra", None) or ""),
        ("Fan RLA", lambda d, c: getattr(d, "fan_rla", None) or ""),
        ("Fan Count", lambda d, c: getattr(d, "fan_count", None) or ""),
        ("Evaporator Count", lambda d, c: getattr(d, "evaporator_count", None) or ""),
        ("SEER", lambda d, c: getattr(d, "seer", None) or ""),
        ("EER", lambda d, c: getattr(d, "eer", None) or ""),
        ("IPLV", lambda d, c: getattr(d, "iplv", None) or ""),
        ("AHRI #", lambda d, c: getattr(d, "ahri_number", None) or ""),
        ("Notes", lambda d, c: getattr(d, "notes", None) or ""),
    ]

    decorated = []
    for d in devices:
        c = classify_system({
            "category": getattr(d, "category", None) or (d.get("category") if isinstance(d, dict) else None),
            "subcategory": getattr(d, "subcategory", None) or (d.get("subcategory") if isinstance(d, dict) else None),
            "equipment_type": getattr(d, "equipment_type", None) or (d.get("equipment_type") if isinstance(d, dict) else None),
            "device_name": getattr(d, "device_name", None) or (d.get("device_name") if isinstance(d, dict) else None),
            "name": getattr(d, "name", None) or (d.get("name") if isinstance(d, dict) else None),
            "model_number": getattr(d, "model_number", None) or (d.get("model_number") if isinstance(d, dict) else None),
            "manufacturer": getattr(d, "manufacturer", None) or (d.get("manufacturer") if isinstance(d, dict) else None),
        })
        decorated.append((d, c))

    decorated.sort(
        key=lambda item: (
            GROUP_RANK.get(item[1].get("group"), 9),
            str(item[1].get("system") or ""),
            str(getattr(item[0], "manufacturer", "") or ""),
        )
    )

    header = ",".join(col[0] for col in equipment_columns)
    lines = [header]

    for d, c in decorated:
        row = [csv_cell(extractor(d, c)) for _, extractor in equipment_columns]
        lines.append(",".join(row))

    return "\r\n".join(lines)


def build_equipment_excel(devices: List[Any], sheet_name: str = "Equipment Inventory") -> bytes:
    """Fallback workbook generator that outputs tabular equipment specs."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    equipment_columns = [
        ("Name", lambda d, c: getattr(d, "device_name", None) or getattr(d, "name", None) or (d.get("name") if isinstance(d, dict) else "")),
        ("Type", lambda d, c: c.get("system")),
        ("Category", lambda d, c: GROUP_LABEL.get(c.get("group"), c.get("group"))),
        ("Qty", lambda d, c: getattr(d, "quantity", 1) or 1),
        ("Manufacturer", lambda d, c: getattr(d, "manufacturer", None) or (d.get("manufacturer") if isinstance(d, dict) else "")),
        ("Model", lambda d, c: getattr(d, "model_number", None) or (d.get("model_number") if isinstance(d, dict) else "")),
        ("Serial", lambda d, c: getattr(d, "serial_number", None) or (d.get("serial_number") if isinstance(d, dict) else "")),
        ("Age (yrs)", lambda d, c: calc_age(getattr(d, "mfg_year", None), getattr(d, "mfg_date", None), getattr(d, "serial_number", None), getattr(d, "manufacturer", None))),
        ("Mfg Date", lambda d, c: getattr(d, "mfg_date", None) or (str(getattr(d, "mfg_year", "")) if getattr(d, "mfg_year", None) else "")),
        ("Tonnage", lambda d, c: getattr(d, "tonnage", None) or (d.get("tonnage") if isinstance(d, dict) else "")),
        ("BTU", lambda d, c: getattr(d, "btu", None) or ""),
        ("Voltage", lambda d, c: getattr(d, "voltage", None) or (d.get("voltage") if isinstance(d, dict) else "")),
        ("Compressor PH", lambda d, c: getattr(d, "compressor_ph", None) or ""),
        ("Fan PH", lambda d, c: getattr(d, "fan_phases", None) or getattr(d, "fan_ph", None) or getattr(d, "compressor_ph", None) or (d.get("fan_phases") if isinstance(d, dict) else (d.get("fan_ph") if isinstance(d, dict) else (d.get("compressor_ph") if isinstance(d, dict) else (d.get("phase") if isinstance(d, dict) else "")))) or ""),
        ("Refrigerant", lambda d, c: getattr(d, "refrigerant_type", None) or (d.get("refrigerant") if isinstance(d, dict) else "")),
        ("MCA", lambda d, c: getattr(d, "mca", None) or ""),
        ("MOCP", lambda d, c: getattr(d, "mocp", None) or ""),
        ("Compressor RLA", lambda d, c: getattr(d, "compressor_rla", None) or ""),
        ("Compressor LRA", lambda d, c: getattr(d, "compressor_lra", None) or ""),
        ("Fan RLA", lambda d, c: getattr(d, "fan_rla", None) or ""),
        ("Fan Count", lambda d, c: getattr(d, "fan_count", None) or ""),
        ("Evaporator Count", lambda d, c: getattr(d, "evaporator_count", None) or ""),
        ("SEER", lambda d, c: getattr(d, "seer", None) or ""),
        ("EER", lambda d, c: getattr(d, "eer", None) or ""),
        ("IPLV", lambda d, c: getattr(d, "iplv", None) or ""),
        ("AHRI #", lambda d, c: getattr(d, "ahri_number", None) or ""),
        ("Notes", lambda d, c: getattr(d, "notes", None) or ""),
    ]

    headers = [col[0] for col in equipment_columns]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A2332", end_color="1A2332", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    decorated = []
    for d in devices:
        c = classify_system({
            "category": getattr(d, "category", None) or (d.get("category") if isinstance(d, dict) else None),
            "subcategory": getattr(d, "subcategory", None) or (d.get("subcategory") if isinstance(d, dict) else None),
            "equipment_type": getattr(d, "equipment_type", None) or (d.get("equipment_type") if isinstance(d, dict) else None),
            "device_name": getattr(d, "device_name", None) or (d.get("device_name") if isinstance(d, dict) else None),
            "name": getattr(d, "name", None) or (d.get("name") if isinstance(d, dict) else None),
            "model_number": getattr(d, "model_number", None) or (d.get("model_number") if isinstance(d, dict) else None),
            "manufacturer": getattr(d, "manufacturer", None) or (d.get("manufacturer") if isinstance(d, dict) else None),
        })
        decorated.append((d, c))

    decorated.sort(
        key=lambda item: (
            GROUP_RANK.get(item[1].get("group"), 9),
            str(item[1].get("system") or ""),
            str(getattr(item[0], "manufacturer", "") or ""),
        )
    )

    for d, c in decorated:
        row_vals = [extractor(d, c) for _, extractor in equipment_columns]
        ws.append([v if v is not None else "" for v in row_vals])

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_submission_excel(job: Any, equipment_items: List[Any], billing_records: Optional[List[Any]] = None) -> bytes:
    """Populates the official Madison Submission Form 8 Excel template with PostgreSQL data."""
    import openpyxl

    template_file = None
    for candidate in TEMPLATE_CANDIDATES:
        if os.path.exists(candidate):
            template_file = candidate
            break

    if not template_file:
        logger.warning("Submission Form template not found in candidates %s; falling back to tabular build.", TEMPLATE_CANDIDATES)
        safe_name = getattr(job, "client_name", None) or getattr(job, "name", None) or "Equipment"
        return build_equipment_excel(equipment_items, sheet_name=str(safe_name)[:31])

    wb = openpyxl.load_workbook(template_file, data_only=False)

    # =========================================================================
    # 1. Populate Input Form
    # =========================================================================
    if "Input Form" in wb.sheetnames:
        ws_input = wb["Input Form"]

        # --- A. Project Information ---
        job_name = getattr(job, "name", None) or getattr(job, "client_name", None) or ""
        ws_input["C10"] = job_name
        ws_input["C9"] = job_name
        ws_input["C11"] = getattr(job, "address", "") or ""
        ws_input["C12"] = getattr(job, "city", "") or ""
        ws_input["C13"] = getattr(job, "state", "") or ""
        ws_input["C14"] = getattr(job, "building_type", "") or ""
        ws_input["C15"] = getattr(job, "entity_type", "") or ""
        ws_input["C16"] = getattr(job, "utility_name", "") or getattr(job, "billing_utility_name", "") or ""

        sqft = getattr(job, "sqft", None)
        if sqft is not None:
            try:
                ws_input["C17"] = float(sqft)
            except (ValueError, TypeError):
                ws_input["C17"] = str(sqft)

        hours_day = getattr(job, "operating_hours_per_day", None)
        if hours_day is not None:
            try:
                ws_input["C19"] = float(hours_day)
            except (ValueError, TypeError):
                pass

        days_week = getattr(job, "operating_days_per_week", None)
        if days_week is not None:
            try:
                ws_input["D19"] = float(days_week)
            except (ValueError, TypeError):
                pass

        ws_input["C25"] = getattr(job, "contact_name", "") or ""
        ws_input["C26"] = getattr(job, "contact_email", "") or ""
        ws_input["C27"] = getattr(job, "contact_phone", "") or ""
        ws_input["H13"] = getattr(job, "advisor", "") or ""

        notes = getattr(job, "notes", "") or ""
        if notes:
            ws_input["E49"] = notes

        # --- B. Monthly Energy Billing Data (Rows 33 to 44) ---
        billings = billing_records or (list(getattr(job, "billing_records", [])) if hasattr(job, "billing_records") else [])
        for b in billings:
            m_str = str(getattr(b, "month", "")).strip().lower()
            row_idx = MONTH_ROW_MAP.get(m_str)
            if row_idx:
                kwh = getattr(b, "kwh_usage", None)
                demand = getattr(b, "demand_kw", None)
                energy_chg = getattr(b, "energy_charge", None)
                demand_chg = getattr(b, "demand_charge", None)
                taxes = getattr(b, "taxes", None)

                if kwh is not None:
                    ws_input.cell(row=row_idx, column=6).value = float(kwh)
                if demand is not None:
                    ws_input.cell(row=row_idx, column=7).value = float(demand)
                if energy_chg is not None:
                    ws_input.cell(row=row_idx, column=8).value = float(energy_chg)
                if demand_chg is not None:
                    ws_input.cell(row=row_idx, column=9).value = float(demand_chg)
                if taxes is not None:
                    ws_input.cell(row=row_idx, column=10).value = float(taxes)

        # --- C. HVAC & Cold Storage Equipment Tables (Rows 58 to 193) ---
        section_rows = {
            "RTU_460": {"curr": 58, "max": 73},
            "RTU_230": {"curr": 79, "max": 93},
            "SPLIT":   {"curr": 100, "max": 114},
            "CHILLER": {"curr": 122, "max": 126},
            "PTAC":    {"curr": 133, "max": 137},
            "VTAC":    {"curr": 144, "max": 148},
            "WIC":     {"curr": 155, "max": 164},
            "WIF":     {"curr": 168, "max": 177},
            "FBC":     {"curr": 181, "max": 185},
            "FBF":     {"curr": 189, "max": 193},
        }

        for item in equipment_items:
            analysis = getattr(item, "analysis_result", None) or {}
            mfg = (getattr(item, "manufacturer", None) or analysis.get("manufacturer") or "").strip()
            model = (getattr(item, "model_number", None) or analysis.get("model_number") or "").strip()
            serial = (getattr(item, "serial_number", None) or analysis.get("serial_number") or "").strip()
            eq_type = (getattr(item, "equipment_type", None) or analysis.get("equipment_type") or "").strip()
            dev_name = (getattr(item, "device_name", None) or getattr(item, "nameplate_label", None) or getattr(item, "name", None) or "").strip()

            c = classify_system({
                "category": getattr(item, "category", None),
                "subcategory": getattr(item, "subcategory", None),
                "equipment_type": eq_type,
                "device_name": dev_name,
                "name": getattr(item, "name", None),
                "model_number": model,
                "manufacturer": mfg,
            })
            sys_type = c.get("system", "")
            group = c.get("group", "")
            volt_str = str(getattr(item, "voltage", None) or analysis.get("voltage") or "").lower()

            sec_key = None
            if "chiller" in sys_type.lower():
                sec_key = "CHILLER"
            elif "ptac" in sys_type.lower():
                sec_key = "PTAC"
            elif "vtac" in sys_type.lower():
                sec_key = "VTAC"
            elif "split" in sys_type.lower() or sys_type in ["Heat Pump", "VRF/VRV", "Air Handler", "Fan Coil", "Furnace"]:
                sec_key = "SPLIT"
            elif "walk" in sys_type.lower() and "freezer" in sys_type.lower():
                sec_key = "WIF"
            elif "walk" in sys_type.lower() or "cooler" in sys_type.lower():
                sec_key = "WIC"
            elif "f&b" in sys_type.lower() and "freezer" in sys_type.lower():
                sec_key = "FBF"
            elif "f&b" in sys_type.lower() or "beverage" in sys_type.lower() or "display" in sys_type.lower():
                sec_key = "FBC"
            elif group == "cold":
                sec_key = "WIC"
            else:  # RTU or general HVAC
                if any(v in volt_str for v in ["460", "480", "575"]):
                    sec_key = "RTU_460"
                else:
                    sec_key = "RTU_230"

            sec = section_rows.get(sec_key)
            if not sec or sec["curr"] > sec["max"]:
                if sec_key != "RTU_460" and section_rows["RTU_460"]["curr"] <= section_rows["RTU_460"]["max"]:
                    sec_key = "RTU_460"
                elif sec_key != "RTU_230" and section_rows["RTU_230"]["curr"] <= section_rows["RTU_230"]["max"]:
                    sec_key = "RTU_230"
                elif section_rows["SPLIT"]["curr"] <= section_rows["SPLIT"]["max"]:
                    sec_key = "SPLIT"
                sec = section_rows[sec_key]

            if sec and sec["curr"] <= sec["max"]:
                r = sec["curr"]
                sec["curr"] += 1

                mfg_year = getattr(item, "mfg_year", None)
                mfg_date = getattr(item, "mfg_date", None)
                age = calc_age(mfg_year, mfg_date, serial, mfg)
                age_num = clean_int(age)
                tonnage_num = clean_float(getattr(item, "tonnage", None) or analysis.get("tonnage"))
                mca_num = clean_float(getattr(item, "mca", None) or analysis.get("mca") or analysis.get("min_circuit_ampacity"))
                rla_num = clean_float(getattr(item, "compressor_rla", None) or analysis.get("compressor_rla"))
                fan_ph_val = getattr(item, "fan_phases", None) or getattr(item, "fan_ph", None) or analysis.get("fan_phases") or analysis.get("fan_ph") or getattr(item, "compressor_ph", None) or analysis.get("compressor_ph") or analysis.get("phase") or getattr(item, "phase", None)
                voltage_val = getattr(item, "voltage", None) or analysis.get("voltage")

                if sec_key in ["RTU_460", "RTU_230", "SPLIT"]:
                    ws_input.cell(r, 3).value = getattr(item, "quantity", 1) or 1
                    ws_input.cell(r, 4).value = model or ""
                    ws_input.cell(r, 5).value = serial or ""
                    if tonnage_num: ws_input.cell(r, 6).value = tonnage_num
                    ws_input.cell(r, 7).value = 460 if sec_key == "RTU_460" else (230 if sec_key == "RTU_230" else (voltage_val or ""))
                    if fan_ph_val: ws_input.cell(r, 8).value = clean_int(fan_ph_val)
                    if mca_num: ws_input.cell(r, 9).value = mca_num
                    if rla_num: ws_input.cell(r, 11).value = rla_num
                    if age_num: ws_input.cell(r, 12).value = age_num

                elif sec_key == "CHILLER":
                    ws_input.cell(r, 3).value = getattr(item, "quantity", 1) or 1
                    ws_input.cell(r, 4).value = model or ""
                    ws_input.cell(r, 5).value = serial or ""
                    if tonnage_num: ws_input.cell(r, 6).value = tonnage_num
                    ws_input.cell(r, 7).value = voltage_val or 460
                    if fan_ph_val: ws_input.cell(r, 8).value = clean_int(fan_ph_val)
                    if mca_num: ws_input.cell(r, 9).value = mca_num
                    comp_count = getattr(item, "compressor_count", None) or analysis.get("compressor_count")
                    if comp_count: ws_input.cell(r, 10).value = comp_count
                    if rla_num: ws_input.cell(r, 11).value = rla_num
                    if age_num: ws_input.cell(r, 12).value = age_num

                elif sec_key in ["PTAC", "VTAC"]:
                    ws_input.cell(r, 3).value = getattr(item, "quantity", 1) or 1
                    ws_input.cell(r, 4).value = model or ""
                    ws_input.cell(r, 5).value = serial or ""
                    if tonnage_num: ws_input.cell(r, 6).value = tonnage_num
                    ws_input.cell(r, 7).value = voltage_val or 230
                    if fan_ph_val: ws_input.cell(r, 8).value = clean_int(fan_ph_val)
                    if mca_num: ws_input.cell(r, 9).value = mca_num
                    fan_rla_val = getattr(item, "fan_rla", None) or analysis.get("fan_rla")
                    if fan_rla_val: ws_input.cell(r, 10).value = clean_float(fan_rla_val)
                    if rla_num: ws_input.cell(r, 11).value = rla_num
                    if age_num: ws_input.cell(r, 12).value = age_num

                elif sec_key in ["WIC", "WIF", "FBC", "FBF"]:
                    evap_cnt = getattr(item, "evaporator_count", None) or analysis.get("evaporator_count")
                    if evap_cnt: ws_input.cell(r, 3).value = evap_cnt
                    fan_cnt = getattr(item, "fan_count", None) or analysis.get("fan_count")
                    if fan_cnt: ws_input.cell(r, 4).value = fan_cnt
                    ws_input.cell(r, 5).value = model or ""
                    ws_input.cell(r, 6).value = serial or ""
                    ws_input.cell(r, 7).value = voltage_val or 230
                    if tonnage_num: ws_input.cell(r, 8).value = tonnage_num
                    comp_hp = getattr(item, "compressor_hp", None) or analysis.get("compressor_hp")
                    if comp_hp: ws_input.cell(r, 9).value = clean_float(comp_hp)
                    mot_type = getattr(item, "motor_type", None) or analysis.get("motor_type")
                    if mot_type: ws_input.cell(r, 10).value = mot_type
                    fan_rla_val = getattr(item, "fan_rla", None) or analysis.get("fan_rla") or getattr(item, "fan_fla", None)
                    if fan_rla_val: ws_input.cell(r, 11).value = clean_float(fan_rla_val)
                    if age_num: ws_input.cell(r, 12).value = age_num

    # =========================================================================
    # 2. Populate Panel Detail (Distribution Panels & Breakers)
    # =========================================================================
    if "Panel Detail" in wb.sheetnames:
        ws_panel = wb["Panel Detail"]

        # Clear default sample values from template
        for sample_cell in ["B10", "B11", "B12", "H10", "H11", "H12", "N10", "N11", "N12", "T10", "T11", "T12"]:
            if ws_panel[sample_cell].value == "RTU1":
                ws_panel[sample_cell].value = None

        # Clear [insert panel name] placeholder text
        for p in PANEL_SPECS:
            cell = ws_panel[p["name_cell"]]
            if cell.value and "[insert panel name]" in str(cell.value):
                cell.value = None

        panels_used = set()
        for idx, item in enumerate(equipment_items):
            panel_idx = idx // 44
            if panel_idx >= len(PANEL_SPECS):
                break

            p = PANEL_SPECS[panel_idx]
            panels_used.add(p["num"])

            slot = idx % 44
            if slot < 22:
                col = p["left_col"]
                row = p["start_row"] + slot
            else:
                col = p["right_col"]
                row = p["start_row"] + (slot - 22)

            label = format_equipment_item(item, idx)
            ws_panel.cell(row=row, column=col).value = label

        if not panels_used:
            panels_used.add(1)

        for p_num in sorted(panels_used):
            p = PANEL_SPECS[p_num - 1]
            ws_panel[p["name_cell"]].value = f"Panel {p_num}"
            if "Input Form" in wb.sheetnames and p_num <= 12:
                ws_input = wb["Input Form"]
                ws_input.cell(row=32 + p_num, column=2).value = f"Panel {p_num}"
                if ws_input.cell(row=32 + p_num, column=3).value is None:
                    ws_input.cell(row=32 + p_num, column=3).value = 0

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@router.get("/{job_id}/export/equipment.csv")
async def export_equipment_csv(
    job_id: uuid.UUID,
    auth: AuthContext = Depends(get_auth_context),
    session: AsyncSession = Depends(get_session),
):
    job = await get_job_by_id(session, job_id, auth.org_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    devices = await list_devices_for_job(session, job_id, auth.org_id)
    csv_data = build_equipment_csv(devices)

    safe_name = re.sub(r"[^a-zA-Z0-9\s\-_]", "", job.name or job.client_name or job.job_number or str(job_id))
    safe_name = re.sub(r"\s+", "_", safe_name.strip()) or "Site_Hunter_Export"
    filename = f"{safe_name}_Equipment_Inventory.csv"

    return Response(
        content=csv_data,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@router.get("/{job_id}/export/excel")
async def export_submission_excel(
    job_id: uuid.UUID,
    auth: AuthContext = Depends(get_auth_context),
    session: AsyncSession = Depends(get_session),
):
    job = await get_job_by_id(session, job_id, auth.org_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    job_name = getattr(job, "name", None) or getattr(job, "client_name", None) or str(job_id)
    setattr(job, "name", job_name)

    devices = list(getattr(job, "devices", [])) or await list_devices_for_job(session, job_id, auth.org_id)
    scans = []
    for d in devices:
        if hasattr(d, "scans") and d.scans:
            scans.extend(d.scans)

    job.scans = scans if scans else devices
    logger.info(f"Exporting job: {job.name} with {len(job.scans)} scans")

    equipment_items = devices if len(devices) >= len(job.scans) else job.scans
    billings = list(getattr(job, "billing_records", []))

    safe_name = re.sub(r"[^a-zA-Z0-9\s\-_]", "", job.name)
    safe_name = re.sub(r"\s+", "_", safe_name.strip()) or "Site_Hunter_Export"
    filename = f"{safe_name}_Project_Submission_Form.xlsx"

    excel_bytes = build_submission_excel(job, equipment_items, billings)
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


async def handle_export_excel(job_id_or_gid: str, request: Optional[Any] = None) -> Response:
    """Direct handler for /api/export/excel/{job_id} supporting UUIDs, Firestore IDs, Asana GIDs, and Job Numbers."""
    from app.db import async_session_factory as AsyncSessionLocal
    from app.db_models import Job, Device, Scan
    from sqlalchemy import select, or_
    from sqlalchemy.orm import selectinload

    job = None
    async with AsyncSessionLocal() as session:
        # 1. Try UUID lookup
        try:
            job_uuid = uuid.UUID(job_id_or_gid)
            res = await session.execute(
                select(Job).where(Job.id == job_uuid).options(
                    selectinload(Job.devices).selectinload(Device.scans),
                    selectinload(Job.billing_records),
                )
            )
            job = res.scalars().first()
        except (ValueError, TypeError):
            pass

        # 2. Try Firestore ID, Asana GID, or Job Number lookup
        if not job:
            res = await session.execute(
                select(Job).where(
                    or_(
                        Job.firestore_id == str(job_id_or_gid),
                        Job.asana_project_gid == str(job_id_or_gid),
                        Job.job_number == str(job_id_or_gid),
                    )
                ).options(
                    selectinload(Job.devices).selectinload(Device.scans),
                    selectinload(Job.billing_records),
                )
            )
            job = res.scalars().first()

        # 3. Try lookup via Scan records if the identifier belongs to a scan
        if not job:
            res_sc = await session.execute(
                select(Scan).where(
                    or_(
                        Scan.firestore_id == str(job_id_or_gid),
                        Scan.job_number == str(job_id_or_gid),
                        Scan.asana_project_gid == str(job_id_or_gid),
                    )
                )
            )
            matched_scan = res_sc.scalars().first()
            if matched_scan:
                res = await session.execute(
                    select(Job).where(
                        or_(
                            Job.job_number == matched_scan.job_number,
                            Job.asana_project_gid == matched_scan.asana_project_gid,
                            Job.firestore_id == matched_scan.job_number,
                        )
                    ).options(
                        selectinload(Job.devices).selectinload(Device.scans),
                        selectinload(Job.billing_records),
                    )
                )
                job = res.scalars().first()

        if not job:
            logger.error(f"Job not found for ID/GID/Firestore ID: '{job_id_or_gid}'")
            raise HTTPException(status_code=404, detail=f"Job not found for ID '{job_id_or_gid}'")

        # Ensure job.name is available
        job_name = getattr(job, "name", None) or getattr(job, "client_name", None) or str(job.id)
        setattr(job, "name", job_name)

        # Collect scans and devices
        devices = list(getattr(job, "devices", []))
        scans = []
        for d in devices:
            if hasattr(d, "scans") and d.scans:
                scans.extend(d.scans)

        res_direct_scans = await session.execute(
            select(Scan).where(
                or_(
                    Scan.job_number == job.job_number,
                    Scan.job_number == str(job.id),
                    Scan.job_number == job.firestore_id,
                    Scan.asana_project_gid == job.asana_project_gid,
                )
            )
        )
        for sc in res_direct_scans.scalars().all():
            if sc not in scans:
                scans.append(sc)

        job.scans = scans if scans else devices
        equipment_items = devices if len(devices) >= len(job.scans) else job.scans
        billings = list(getattr(job, "billing_records", []))

        # Explicit required logger statement
        logger.info(f"Exporting job: {job.name} with {len(job.scans)} scans")

        safe_name = re.sub(r"[^a-zA-Z0-9\s\-_]", "", job.name)
        safe_name = re.sub(r"\s+", "_", safe_name.strip()) or "Site_Hunter_Export"
        filename = f"{safe_name}_Project_Submission_Form.xlsx"

        excel_bytes = build_submission_excel(job, equipment_items, billings)

        return Response(
            content=excel_bytes,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition",
            },
        )


async def handle_export_csv(job_id_or_gid: str, request: Optional[Any] = None) -> Response:
    """Direct handler for /api/export/equipment/{job_id} supporting UUIDs, Firestore IDs, Asana GIDs, and Job Numbers."""
    from app.db import async_session_factory as AsyncSessionLocal
    from app.db_models import Job, Device, Scan
    from sqlalchemy import select, or_
    from sqlalchemy.orm import selectinload

    job = None
    devices = []
    async with AsyncSessionLocal() as session:
        try:
            job_uuid = uuid.UUID(job_id_or_gid)
            res = await session.execute(
                select(Job).where(Job.id == job_uuid).options(selectinload(Job.devices))
            )
            job = res.scalars().first()
            if job:
                devices = list(job.devices)
        except (ValueError, TypeError):
            pass

        if not job:
            res = await session.execute(
                select(Job).where(
                    or_(
                        Job.firestore_id == str(job_id_or_gid),
                        Job.asana_project_gid == str(job_id_or_gid),
                        Job.job_number == str(job_id_or_gid),
                    )
                ).options(selectinload(Job.devices))
            )
            job = res.scalars().first()
            if job:
                devices = list(job.devices)

        if not devices and job:
            res_scans = await session.execute(
                select(Scan).where(
                    or_(
                        Scan.job_number == job.job_number,
                        Scan.job_number == str(job.id),
                        Scan.job_number == job.firestore_id,
                        Scan.asana_project_gid == job.asana_project_gid,
                    )
                )
            )
            devices = list(res_scans.scalars().all())

    job_name = getattr(job, "name", None) or getattr(job, "client_name", None) if job else f"Project_{job_id_or_gid}"
    safe_name = re.sub(r"[^a-zA-Z0-9\s\-_]", "", job_name)
    safe_name = re.sub(r"\s+", "_", safe_name.strip()) or "Site_Hunter_Export"
    filename = f"{safe_name}_Equipment_Inventory.csv"

    csv_data = build_equipment_csv(devices)

    return Response(
        content=csv_data,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


# -----------------------------------------------------------------------------
# Dedicated Project Sync / Calibrator Endpoint
# -----------------------------------------------------------------------------
@router.post("/{job_id}/sync")
@router.post("/sync/{job_id}")
async def export_sync_endpoint(job_id: str, request: Request):
    """Publish a project_sync event for the job to events:calibrator Redis stream."""
    from app.services.calibrator_sync import handle_project_sync_request
    return await handle_project_sync_request(job_id, request=request)


async def handle_export_sync(job_id_or_gid: str, request: Optional[Request] = None) -> Response:
    """Direct handler for /api/export/sync/{job_id} supporting SSE and JSON."""
    from app.services.calibrator_sync import handle_project_sync_request
    return await handle_project_sync_request(job_id_or_gid, request=request)
